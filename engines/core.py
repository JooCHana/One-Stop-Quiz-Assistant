from __future__ import annotations

import csv
import re
import time
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Iterable, Sequence
from xml.etree import ElementTree

from PIL import Image, ImageGrab
from rapidfuzz import fuzz


QUESTION_HEADERS = ("试题内容", "题目", "题干", "问题")
ANSWER_HEADERS = ("正确选项", "正确答案", "答案", "答案1")
TYPE_HEADERS = ("试题类型", "题型", "类型")
ID_HEADERS = ("序号", "编号", "题号")
OPTION_LETTERS = "ABCDEFGH"
QUESTION_TYPES = ("不定项选择题", "多选题", "单选题", "判断题")


class QuizError(RuntimeError):
    """A user-facing error that should stop automation safely."""


@dataclass(frozen=True)
class Question:
    qid: str
    text: str
    qtype: str
    answer: tuple[str, ...]
    options: dict[str, str]
    source_sheet: str = ""
    source_row: int = 0
    normalized_text: str = field(init=False, repr=False)

    def __post_init__(self) -> None:
        object.__setattr__(self, "normalized_text", normalize_text(self.text))


@dataclass(frozen=True)
class OCRItem:
    box: tuple[tuple[float, float], ...]
    text: str
    confidence: float

    @property
    def center(self) -> tuple[float, float]:
        xs = [p[0] for p in self.box]
        ys = [p[1] for p in self.box]
        return sum(xs) / len(xs), sum(ys) / len(ys)

    @property
    def width(self) -> float:
        xs = [p[0] for p in self.box]
        return max(xs) - min(xs)

    @property
    def height(self) -> float:
        ys = [p[1] for p in self.box]
        return max(ys) - min(ys)

    @property
    def horizontal(self) -> bool:
        if len(self.box) < 2:
            return True
        return abs(self.box[0][1] - self.box[1][1]) <= max(10.0, self.height * 0.35)


@dataclass
class PageAnalysis:
    question: Question
    score: float
    second_score: float
    recognized_question: str
    items: list[OCRItem]
    image_size: tuple[int, int]
    question_y: float
    progress: tuple[int, int] | None
    answer_points: dict[str, tuple[int, int]]
    next_point: tuple[int, int] | None
    finish_button_detected: bool

    @property
    def page_key(self) -> str:
        if self.progress:
            return f"{self.progress[0]}/{self.progress[1]}"
        return self.question.normalized_text


@dataclass
class AnswerView:
    items: list[OCRItem]
    answer_points: dict[str, tuple[int, int]]
    progress: tuple[int, int] | None
    next_point: tuple[int, int] | None
    finish_button_detected: bool


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text or "").lower()
    return "".join(ch for ch in text if ch.isalnum())


def clean_question_text(text: str) -> str:
    value = unicodedata.normalize("NFKC", text or "")
    for token in QUESTION_TYPES:
        value = value.replace(token, "")
    value = re.sub(r"^\s*(?:第\s*)?\d+\s*[、.．:]\s*", "", value)
    value = re.sub(r"[（(]\s*\d+(?:\.\d+)?\s*分\s*[）)]\s*$", "", value)
    return value.strip()


def normalize_answer(value: str) -> tuple[str, ...]:
    letters = re.findall(r"[A-H]", unicodedata.normalize("NFKC", value or "").upper())
    return tuple(dict.fromkeys(letters))


def _find_header_row(rows: Sequence[Sequence[object]]) -> tuple[int, dict[str, int]]:
    for row_idx, row in enumerate(rows[:40]):
        headers = {cell_text(v).replace("\n", "").replace(" ", ""): i for i, v in enumerate(row)}
        q_idx = next((headers[h] for h in QUESTION_HEADERS if h in headers), None)
        a_idx = next((headers[h] for h in ANSWER_HEADERS if h in headers), None)
        if q_idx is None or a_idx is None:
            continue
        mapping = {"question": q_idx, "answer": a_idx}
        for key, names in (("type", TYPE_HEADERS), ("id", ID_HEADERS)):
            idx = next((headers[h] for h in names if h in headers), None)
            if idx is not None:
                mapping[key] = idx
        for letter in OPTION_LETTERS:
            number = ord(letter) - ord('A') + 1
            aliases = (letter, f"选项{letter}", f"选择项{number}", f"选项{number}")
            idx = next((headers[name] for name in aliases if name in headers), None)
            if idx is not None:
                mapping[letter] = idx
        return row_idx, mapping
    raise QuizError("未找到题库表头。至少需要“试题内容”和“正确选项/答案”两列。")


def _questions_from_rows(
    rows: Sequence[Sequence[object]], sheet_name: str, log: Callable[[str], None] | None = None
) -> list[Question]:
    if not rows:
        return []
    try:
        header_row, columns = _find_header_row(rows)
    except QuizError:
        return []

    questions: list[Question] = []
    for row_idx in range(header_row + 1, len(rows)):
        row = rows[row_idx]

        def get(name: str) -> str:
            idx = columns.get(name)
            return cell_text(row[idx]) if idx is not None and idx < len(row) else ""

        text = get("question")
        answer = normalize_answer(get("answer"))
        if not text or not answer:
            continue
        qtype = get("type")
        qtype = {"判断": "判断题", "单选": "单选题", "多选": "多选题"}.get(qtype, qtype)
        if not qtype:
            qtype = "多选题" if len(answer) > 1 else "单选题"
        if qtype and not any(token in qtype for token in QUESTION_TYPES):
            continue
        options = {letter: get(letter) for letter in OPTION_LETTERS if get(letter)}
        if "判断题" in qtype and not options:
            options = {"A": "正确", "B": "错误"}
        if not options:
            continue
        valid_answer = tuple(letter for letter in answer if letter in options)
        if not valid_answer:
            continue
        questions.append(
            Question(
                qid=get("id") or str(row_idx + 1),
                text=text,
                qtype=qtype,
                answer=valid_answer,
                options=options,
                source_sheet=sheet_name,
                source_row=row_idx + 1,
            )
        )
    if log and questions:
        log(f"工作表“{sheet_name}”读取 {len(questions)} 道选择题")
    return questions


def _extract_docx_text(source: Path) -> str:
    try:
        with zipfile.ZipFile(source) as archive:
            document_xml = archive.read("word/document.xml")
    except (OSError, KeyError, zipfile.BadZipFile) as exc:
        raise QuizError("Word 文档无法读取或文件已经损坏。") from exc
    try:
        root = ElementTree.fromstring(document_xml)
    except ElementTree.ParseError as exc:
        raise QuizError("Word 文档正文格式无法解析。") from exc
    parts: list[str] = []
    for element in root.iter():
        name = element.tag.rsplit("}", 1)[-1]
        if name == "t" and element.text:
            parts.append(element.text)
        elif name in ("tab", "br"):
            parts.append(" ")
        elif name in ("p", "tr"):
            parts.append("\n")
    return "".join(parts)


def extract_document_text(source: Path) -> str:
    suffix = source.suffix.lower()
    if suffix == ".doc":
        try:
            from legacy_doc import extract_text
        except ImportError as exc:
            raise QuizError("缺少旧版 Word 文档读取组件，请重新安装完整版本软件。") from exc
        try:
            return extract_text(source.read_bytes()).text
        except Exception as exc:
            raise QuizError(
                "旧版 Word 文档无法读取。文件可能已损坏、加密，或不是标准 Word 97-2003 格式。"
            ) from exc
    if suffix == ".docx":
        return _extract_docx_text(source)
    if suffix == ".txt":
        for encoding in ("utf-8-sig", "gb18030", "utf-16"):
            try:
                return source.read_text(encoding=encoding)
            except UnicodeError:
                continue
        raise QuizError("TXT 文档编码无法识别，请另存为 UTF-8 或 GB18030。")
    raise QuizError("当前文档格式不受支持。")


def _document_section(text: str, keyword: str) -> str:
    heading = re.compile(rf"[一二三四五六七八九十]+\s*[、.．]\s*{keyword}")
    match = heading.search(text)
    if not match:
        return ""
    next_heading = re.compile(r"[一二三四五六七八九十]+\s*[、.．]\s*(?:填空题|选择题|判断题|简答题?|问答题)")
    following = next_heading.search(text, match.end())
    return text[match.end() : following.start() if following else len(text)]


def _numbered_document_items(section: str) -> list[tuple[str, str]]:
    pattern = re.compile(r"(?m)(?:^|\n)\s*(\d{1,3})\s*[.．、]\s*")
    matches = list(pattern.finditer(section))
    result: list[tuple[str, str]] = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(section)
        body = section[match.end() : end].strip()
        if body:
            result.append((match.group(1), body))
    return result


def _parse_document_choice_item(
    number: str, body: str, source: Path
) -> Question | None:
    value = unicodedata.normalize("NFKC", body).replace("\r", "\n")
    marker_pattern = re.compile(r"(?:^|(?<=\s))([A-H])\s*[、.．:：]\s*", re.MULTILINE)
    markers = list(marker_pattern.finditer(value))
    if len(markers) < 2:
        return None
    question_text = value[: markers[0].start()].strip()
    answer_match = re.search(r"【\s*([A-H](?:[\s、,，/]*[A-H])*)\s*】", question_text, re.I)
    if not answer_match:
        return None
    answer = normalize_answer(answer_match.group(1))
    question_text = (question_text[: answer_match.start()] + question_text[answer_match.end() :]).strip()
    options: dict[str, str] = {}
    for index, marker in enumerate(markers):
        end = markers[index + 1].start() if index + 1 < len(markers) else len(value)
        option = re.sub(r"\s+", " ", value[marker.end() : end]).strip(" ;；")
        if option:
            options[marker.group(1)] = option
    valid_answer = tuple(letter for letter in answer if letter in options)
    if not question_text or not valid_answer:
        return None
    return Question(
        qid=f"选择{number}",
        text=question_text,
        qtype="多选题" if len(valid_answer) > 1 else "单选题",
        answer=valid_answer,
        options=options,
        source_sheet=source.name,
        source_row=int(number),
    )


def _parse_document_judgment_item(
    number: str, body: str, source: Path
) -> Question | None:
    value = unicodedata.normalize("NFKC", body).replace("\r", "\n")
    answer_match = re.search(r"[【\[]\s*([√×xX对错])\s*[】\]]", value)
    if not answer_match:
        return None
    answer_value = answer_match.group(1)
    question_text = (value[: answer_match.start()] + value[answer_match.end() :]).strip()
    question_text = re.sub(r"\s+", " ", question_text)
    if not question_text:
        return None
    return Question(
        qid=f"判断{number}",
        text=question_text,
        qtype="判断题",
        answer=("A",) if answer_value in ("√", "对") else ("B",),
        options={"A": "正确", "B": "错误"},
        source_sheet=source.name,
        source_row=int(number),
    )


def questions_from_document_text(
    text: str,
    source: Path,
    log: Callable[[str], None] | None = None,
) -> list[Question]:
    normalized = unicodedata.normalize("NFKC", text).replace("\x07", " ").replace("\r", "\n")
    choice_items = _numbered_document_items(_document_section(normalized, "选择题"))
    judgment_items = _numbered_document_items(_document_section(normalized, "判断题"))
    fill_count = len(_numbered_document_items(_document_section(normalized, "填空题")))
    short_section = _document_section(normalized, "(?:简答题?|问答题)")
    short_count = sum(
        1
        for line in short_section.splitlines()
        if re.match(r"\s*\d{1,3}\s*[.．、]", line) and ("?" in line or "？" in line)
    )
    questions = [
        question
        for number, body in choice_items
        if (question := _parse_document_choice_item(number, body, source)) is not None
    ]
    questions.extend(
        question
        for number, body in judgment_items
        if (question := _parse_document_judgment_item(number, body, source)) is not None
    )
    if log:
        choice_count = sum(question.qtype != "判断题" for question in questions)
        judgment_count = sum(question.qtype == "判断题" for question in questions)
        log(
            f"文档读取完成：可自动作答的选择题 {choice_count} 道、判断题 {judgment_count} 道"
        )
        if fill_count or short_count:
            log(
                f"另检测到填空题 {fill_count} 道、简答题 {short_count} 道；"
                "这两类没有可点击选项，本版不加入自动作答。"
            )
    return questions


def load_question_bank(path: str | Path, log: Callable[[str], None] | None = None) -> list[Question]:
    source = Path(path)
    if not source.exists():
        raise QuizError(f"题库文件不存在：{source}")
    suffix = source.suffix.lower()
    questions: list[Question] = []

    if suffix == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(str(source))
        for sheet in workbook.sheets():
            rows = [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
            questions.extend(_questions_from_rows(rows, sheet.name, log))
    elif suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        workbook = load_workbook(source, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            questions.extend(_questions_from_rows(rows, sheet.title, log))
        workbook.close()
    elif suffix in (".csv", ".tsv"):
        delimiter = "\t" if suffix == ".tsv" else ","
        decoded = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                decoded = source.read_text(encoding=encoding)
                break
            except UnicodeDecodeError:
                continue
        if decoded is None:
            raise QuizError("CSV/TSV 编码无法识别，请另存为 UTF-8 或 GB18030。")
        rows = list(csv.reader(decoded.splitlines(), delimiter=delimiter))
        questions.extend(_questions_from_rows(rows, source.stem, log))
    elif suffix in (".doc", ".docx", ".txt"):
        questions.extend(questions_from_document_text(extract_document_text(source), source, log))
    else:
        raise QuizError(
            "支持的题库格式为 .xls、.xlsx、.xlsm、.csv、.tsv、.doc、.docx 和 .txt。"
        )

    if not questions:
        raise QuizError("题库中没有找到带答案和选项的选择题。")
    return questions


class OfflineOCR:
    def __init__(self) -> None:
        self._engine = None

    def recognize(self, image: Image.Image) -> list[OCRItem]:
        if self._engine is None:
            from rapidocr_onnxruntime import RapidOCR

            self._engine = RapidOCR()
        result, _ = self._engine(image)
        items: list[OCRItem] = []
        for entry in result or []:
            box, text, confidence = entry
            if not text or float(confidence) < 0.42:
                continue
            items.append(
                OCRItem(
                    box=tuple((float(p[0]), float(p[1])) for p in box),
                    text=str(text).strip(),
                    confidence=float(confidence),
                )
            )
        return sorted(items, key=lambda item: (item.center[1], item.center[0]))


def capture_screen_region(rect: tuple[int, int, int, int]) -> Image.Image:
    left, top, right, bottom = rect
    if right - left < 300 or bottom - top < 300:
        raise QuizError("目标窗口尺寸太小，请先恢复并放大答题窗口。")
    return ImageGrab.grab(bbox=rect, all_screens=True)


def _is_ui_or_noise(text: str) -> bool:
    compact = normalize_text(text)
    if not compact:
        return True
    fixed = (
        "立即交卷",
        "上一题",
        "下一题",
        "显示答案",
        "教育培训",
    )
    if any(normalize_text(word) in compact for word in fixed):
        return True
    if re.fullmatch(r"\d{4}\d{1,2}\d{1,2}.*", compact):
        return True
    if re.fullmatch(r"\d+\d+", compact) and len(compact) <= 5:
        return True
    return False


def _question_candidates(items: list[OCRItem], size: tuple[int, int]) -> list[tuple[str, float]]:
    width, height = size
    bottom_y = min(
        (item.center[1] for item in items if "下一题" in item.text or "上一题" in item.text),
        default=height * 0.91,
    )
    qtype_y = min(
        (item.center[1] for item in items if any(token in item.text for token in QUESTION_TYPES)),
        default=height * 0.18,
    )
    option_a_y = min(
        (
            item.center[1]
            for item in items
            if item.text.strip().upper() == "A"
            and item.center[1] > qtype_y
            and item.center[0] < width * 0.4
        ),
        default=bottom_y,
    )

    eligible: list[OCRItem] = []
    for item in items:
        x, y = item.center
        if y < max(height * 0.12, qtype_y - 45) or y >= bottom_y:
            continue
        if x > width * 0.94 or not item.horizontal:
            continue
        if any(token in item.text for token in QUESTION_TYPES) or _is_ui_or_noise(item.text):
            continue
        if re.fullmatch(r"[A-H]", item.text.strip().upper()):
            continue
        cleaned = clean_question_text(item.text)
        if len(normalize_text(cleaned)) >= 5:
            eligible.append(item)

    candidates: list[tuple[str, float]] = [(clean_question_text(item.text), item.center[1]) for item in eligible]

    question_band = [
        item
        for item in eligible
        if qtype_y - 25 <= item.center[1] < option_a_y - 8 and item.center[0] < width * 0.92
    ]
    if question_band:
        joined = "".join(item.text for item in sorted(question_band, key=lambda i: (i.center[1], i.center[0])))
        if len(normalize_text(joined)) >= 5:
            candidates.append((clean_question_text(joined), sum(i.center[1] for i in question_band) / len(question_band)))

    # Wrapped questions may be split into consecutive OCR boxes. Short windows add
    # robustness without turning the whole page (including options) into one candidate.
    for start in range(len(eligible)):
        for length in (2, 3):
            group = eligible[start : start + length]
            if len(group) != length:
                continue
            if group[-1].center[1] - group[0].center[1] > height * 0.16:
                continue
            joined = "".join(item.text for item in group)
            if len(normalize_text(joined)) >= 8:
                candidates.append((clean_question_text(joined), group[0].center[1]))

    unique: dict[str, tuple[str, float]] = {}
    for text, y in candidates:
        normalized = normalize_text(text)
        if normalized and normalized not in unique:
            unique[normalized] = (text, y)
    return list(unique.values())


def _option_region_items(items: list[OCRItem], question_y: float, size: tuple[int, int]) -> list[OCRItem]:
    width, height = size
    bottom_y = min(
        (item.center[1] for item in items if "下一题" in item.text or "上一题" in item.text),
        default=height * 0.91,
    )
    return [
        item
        for item in items
        if question_y + 18 < item.center[1] < bottom_y - 5
        and item.center[0] < width * 0.92
        and item.horizontal
        and not _is_ui_or_noise(item.text)
    ]


def _option_evidence(question: Question, items: list[OCRItem], question_y: float, size: tuple[int, int]) -> float:
    region = _option_region_items(items, question_y, size)
    scores: list[float] = []
    for option in question.options.values():
        option_norm = normalize_text(option)
        if not option_norm:
            continue
        best = 0.0
        for item in region:
            item_norm = normalize_text(item.text)
            if not item_norm:
                continue
            best = max(best, float(fuzz.ratio(option_norm, item_norm)), float(fuzz.partial_ratio(option_norm, item_norm)))
        scores.append(best)
    return sum(scores) / len(scores) if scores else 0.0


def match_question(
    bank: Sequence[Question], items: list[OCRItem], size: tuple[int, int]
) -> tuple[Question, float, float, str, float]:
    candidates = _question_candidates(items, size)
    if not candidates:
        raise QuizError("没有识别到题干。请确认答题页面完整可见，并把窗口适当放大。")

    ranked: list[tuple[float, float, Question, str, float]] = []
    for question in bank:
        best_question_score = 0.0
        best_candidate = ""
        best_y = size[1] * 0.25
        for candidate, y in candidates:
            candidate_norm = normalize_text(candidate)
            if not candidate_norm:
                continue
            score = max(
                float(fuzz.ratio(question.normalized_text, candidate_norm)),
                float(fuzz.partial_ratio(question.normalized_text, candidate_norm)) * 0.985,
            )
            if score > best_question_score:
                best_question_score = score
                best_candidate = candidate
                best_y = y
        evidence = _option_evidence(question, items, best_y, size)
        rank_score = best_question_score + evidence * 0.08
        ranked.append((rank_score, best_question_score, question, best_candidate, best_y))

    ranked.sort(key=lambda row: row[0], reverse=True)
    best = ranked[0]
    second_question_score = ranked[1][1] if len(ranked) > 1 else 0.0
    return best[2], best[1], second_question_score, best[3], best[4]


def _find_progress(items: Iterable[OCRItem]) -> tuple[int, int] | None:
    for item in items:
        match = re.search(r"(\d+)\s*/\s*(\d+)", unicodedata.normalize("NFKC", item.text))
        if match:
            current, total = int(match.group(1)), int(match.group(2))
            if total > 0 and 0 < current <= total:
                return current, total
    return None


def _find_next_point(items: list[OCRItem], size: tuple[int, int]) -> tuple[int, int] | None:
    width, height = size
    candidates = [item for item in items if "下一题" in item.text and item.center[1] > height * 0.65]
    if candidates:
        x, y = max(candidates, key=lambda item: item.center[0]).center
        return round(x), round(y)
    return None


def _finish_button_detected(items: list[OCRItem], size: tuple[int, int]) -> bool:
    height = size[1]
    return any(
        item.center[1] > height * 0.65
        and any(token in item.text for token in ("完成", "提交", "交卷"))
        for item in items
    )


def _option_text_score(target: str, candidate: str) -> float:
    if not target or not candidate:
        return 0.0
    direct = float(fuzz.ratio(target, candidate))
    partial = float(fuzz.partial_ratio(target, candidate))
    coverage = min(len(target), len(candidate)) / max(len(target), len(candidate))
    # A tiny common fragment must not outrank the complete text of another option.
    return max(direct, partial * (0.55 + 0.45 * coverage))


def _best_option_item(option_text: str, region: list[OCRItem]) -> tuple[OCRItem | None, float, float]:
    target = normalize_text(option_text)
    if not target:
        return None, 0.0, 0.0
    ranked: list[tuple[float, OCRItem]] = []
    for item in region:
        candidate = normalize_text(item.text)
        if not candidate or re.fullmatch(r"[a-h]", candidate):
            continue
        ranked.append((_option_text_score(target, candidate), item))
    if not ranked:
        return None, 0.0, 0.0
    ranked.sort(key=lambda row: row[0], reverse=True)
    best_score, best_item = ranked[0]
    second_score = next(
        (
            score
            for score, item in ranked[1:]
            if abs(item.center[1] - best_item.center[1]) > max(18.0, best_item.height * 0.7)
        ),
        0.0,
    )
    return best_item, best_score, best_score - second_score


def detect_option_markers(
    image: Image.Image | None,
    items: list[OCRItem],
    question_y: float,
    size: tuple[int, int],
) -> list[tuple[int, int]]:
    """Detect the left-side option circles/checkboxes without reading their letters."""
    if image is None:
        return []
    try:
        import cv2
        import numpy as np

        rgb = np.asarray(image.convert("RGB"))
        gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
        width, height = size
        bottom_y = min(
            (item.center[1] for item in items if "下一题" in item.text or "上一题" in item.text),
            default=height * 0.92,
        )
        y0 = int(max(height * 0.07, question_y + 24 if question_y >= 0 else height * 0.07))
        y1 = int(min(height * 0.93, bottom_y - 5))
        x0, x1 = int(width * 0.05), int(width * 0.25)
        if y1 - y0 < 80 or x1 - x0 < 80:
            return []
        crop = cv2.GaussianBlur(gray[y0:y1, x0:x1], (5, 5), 1.2)
        min_dim = min(width, height)
        circles = None
        # The second pass also recognizes thin or partly filled option icons.
        # Column grouping and letter/text anchors below still reject stray circles.
        for sensitivity in (27, 21):
            circles = cv2.HoughCircles(
                crop,
                cv2.HOUGH_GRADIENT,
                dp=1.2,
                minDist=max(24, int(height * 0.035)),
                param1=120,
                param2=sensitivity,
                minRadius=max(7, int(min_dim * 0.007)),
                maxRadius=max(18, int(min_dim * 0.024)),
            )
            if circles is not None:
                break
        if circles is None:
            return []
        raw = [(float(x + x0), float(y + y0), float(r)) for x, y, r in circles[0]]

        label_xs = [
            item.center[0]
            for item in items
            if re.fullmatch(r"[A-H]", item.text.strip().upper())
            and item.center[0] < width * 0.4
            and item.center[1] > y0
            and item.center[1] < y1
        ]
        x_tolerance = max(18.0, min_dim * 0.025)
        if label_xs:
            expected_x = sorted(label_xs)[len(label_xs) // 2]
            chosen = [circle for circle in raw if abs(circle[0] - expected_x) <= x_tolerance]
        else:
            groups: list[list[tuple[float, float, float]]] = []
            for circle in sorted(raw, key=lambda row: row[0]):
                placed = False
                for group in groups:
                    mean_x = sum(row[0] for row in group) / len(group)
                    if abs(circle[0] - mean_x) <= max(12.0, min_dim * 0.012):
                        group.append(circle)
                        placed = True
                        break
                if not placed:
                    groups.append([circle])
            groups.sort(
                key=lambda group: (
                    len(group),
                    -abs(sum(row[0] for row in group) / len(group) - width * 0.11),
                ),
                reverse=True,
            )
            chosen = groups[0] if groups and len(groups[0]) >= 2 else []

        points: list[tuple[int, int]] = []
        for x, y, _radius in sorted(chosen, key=lambda row: row[1]):
            if points and abs(y - points[-1][1]) < max(18.0, height * 0.022):
                continue
            points.append((round(x), round(y)))
        return points
    except Exception:
        return []


def _marker_letter_map(
    question: Question,
    markers: list[tuple[int, int]],
    label_map: dict[str, OCRItem],
    region: list[OCRItem],
    question_y: float,
    size: tuple[int, int],
) -> dict[str, tuple[int, int]]:
    if not markers:
        return {}
    option_letters = [letter for letter in OPTION_LETTERS if letter in question.options]
    tolerance = max(30.0, size[1] * 0.035)
    offsets: list[int] = []

    for letter, item in label_map.items():
        if letter not in option_letters:
            continue
        nearest = min(range(len(markers)), key=lambda idx: abs(markers[idx][1] - item.center[1]))
        if abs(markers[nearest][1] - item.center[1]) <= tolerance:
            offsets.append(option_letters.index(letter) - nearest)

    if not offsets:
        # Strong text matches can anchor a scrolled view even when the letter
        # inside the icon was not recognized.
        for letter in option_letters:
            item, score, margin = _best_option_item(question.options[letter], region)
            if item is None or score < 90.0 or (score < 97.0 and margin < 6.0):
                continue
            nearest = min(range(len(markers)), key=lambda idx: abs(markers[idx][1] - item.center[1]))
            if abs(markers[nearest][1] - item.center[1]) <= tolerance:
                offsets.append(option_letters.index(letter) - nearest)

    if offsets:
        offsets.sort()
        offset = offsets[len(offsets) // 2]
        if any(abs(value - offset) > 1 for value in offsets):
            return {}
    elif question_y >= 0 and len(markers) <= len(option_letters):
        # On the initial, unscrolled question view the first marker is A.
        offset = 0
    else:
        return {}

    mapped: dict[str, tuple[int, int]] = {}
    for marker_index, point in enumerate(markers):
        option_index = marker_index + offset
        if 0 <= option_index < len(option_letters):
            mapped[option_letters[option_index]] = point
    return mapped


def find_visible_answer_points(
    question: Question,
    items: list[OCRItem],
    question_y: float,
    size: tuple[int, int],
    letters: Iterable[str] | None = None,
    image: Image.Image | None = None,
) -> dict[str, tuple[int, int]]:
    width, height = size
    region = _option_region_items(items, question_y, size)
    label_map: dict[str, OCRItem] = {}
    for item in region:
        label = item.text.strip().upper()
        if re.fullmatch(r"[A-H]", label) and item.center[0] < width * 0.4:
            previous = label_map.get(label)
            if previous is None or item.center[1] < previous.center[1]:
                label_map[label] = item

    markers = detect_option_markers(image, items, question_y, size)
    marker_map = _marker_letter_map(question, markers, label_map, region, question_y, size)

    wanted = tuple(letters) if letters is not None else question.answer
    points: dict[str, tuple[int, int]] = {}
    for letter in wanted:
        marker_point = marker_map.get(letter)
        if marker_point is not None:
            points[letter] = marker_point
            continue
        label_item = label_map.get(letter)
        if label_item is not None:
            x, y = label_item.center
            points[letter] = (round(x), round(y))
            continue
        option_norm = normalize_text(question.options.get(letter, ""))
        option_item, score, margin = _best_option_item(question.options.get(letter, ""), region)
        required = 82.0 if len(option_norm) <= 3 else (74.0 if len(option_norm) <= 12 else 66.0)
        if option_item is None or score < required or (score < 96.0 and margin < 4.0):
            continue
        x, y = option_item.center
        points[letter] = (round(x), round(y))

    invalid = [
        letter
        for letter, (x, y) in points.items()
        if not (0 <= x < width and max(-1, question_y) < y < height * 0.94)
    ]
    for letter in invalid:
        points.pop(letter, None)
    return points


def scan_answer_view(
    image: Image.Image,
    question: Question,
    ocr: OfflineOCR,
    letters: Iterable[str] | None = None,
) -> AnswerView:
    """Read answer choices and controls from a view after the question was matched.

    The question itself may already be above the viewport after scrolling, so this
    function deliberately does not try to match the bank again.
    """
    items = ocr.recognize(image)
    if not items:
        raise QuizError("下滑后没有识别到页面文字，已停止以避免误点。")
    return AnswerView(
        items=items,
        answer_points=find_visible_answer_points(
            question, items, -20.0, image.size, letters, image=image
        ),
        progress=_find_progress(items),
        next_point=_find_next_point(items, image.size),
        finish_button_detected=_finish_button_detected(items, image.size),
    )


def analyze_page(
    image: Image.Image,
    bank: Sequence[Question],
    ocr: OfflineOCR,
    threshold: float,
) -> PageAnalysis:
    items = ocr.recognize(image)
    if not items:
        raise QuizError("屏幕识别没有返回文字。请确认答题窗口未被遮挡。")
    question, score, second_score, recognized, question_y = match_question(bank, items, image.size)
    score_gap = score - second_score
    if score < threshold or (score < 92.0 and score_gap < 1.0):
        raise QuizError(
            f"题目匹配不够可靠（最高 {score:.1f}%，次高 {second_score:.1f}%）。"
            "程序已停止，没有点击任何选项。"
        )
    answer_points = find_visible_answer_points(question, items, question_y, image.size, image=image)
    return PageAnalysis(
        question=question,
        score=score,
        second_score=second_score,
        recognized_question=recognized,
        items=items,
        image_size=image.size,
        question_y=question_y,
        progress=_find_progress(items),
        answer_points=answer_points,
        next_point=_find_next_point(items, image.size),
        finish_button_detected=_finish_button_detected(items, image.size),
    )


def wait_for_escape(stop_event, duration: float, escape_pressed: Callable[[], bool]) -> bool:
    deadline = time.monotonic() + duration
    while time.monotonic() < deadline:
        if stop_event.is_set() or escape_pressed():
            stop_event.set()
            return True
        time.sleep(min(0.08, max(0.0, deadline - time.monotonic())))
    return stop_event.is_set()
